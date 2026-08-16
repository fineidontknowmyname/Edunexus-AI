import logging
from pathlib import Path

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".pptx", ".docx", ".txt")


def extract_text(file_path: str, filename: str) -> str:
    print(f"[PARSER] Starting extraction for file: {filename} at path: {file_path}")

    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        print(f"[PARSER ERROR] Unsupported file extension '{extension}' for file '{filename}'")
        raise ValueError(
            f"Unsupported file type '{extension}'. "
            f"Supported formats: {', '.join(SUPPORTED_EXTENSIONS)}"
        )

    path = Path(file_path)
    if not path.exists():
        print(f"[PARSER ERROR] File not found at path: {file_path}")
        raise FileNotFoundError(f"File not found: {file_path}")

    print(f"[PARSER] Detected extension '{extension}'. Routing to handler.")

    try:
        if extension == ".pdf":
            text = _parse_pdf(path)
        elif extension == ".pptx":
            text = _parse_pptx(path)
        elif extension == ".docx":
            text = _parse_docx(path)
        else:
            text = _parse_txt(path)
    except ValueError:
        raise
    except Exception as e:
        print(f"[PARSER ERROR] Unexpected exception while parsing '{filename}': {e}")
        logger.exception("Unexpected error in extract_text for '%s'", filename)
        raise

    print(f"[PARSER SUCCESS] Extracted {len(text):,} characters from '{filename}'")
    return text


def _parse_pdf(path: Path) -> str:
    print(f"[PDF PARSER] Opening PDF with PyMuPDF: {path}")

    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # type: ignore[no-redef]

    try:
        doc = fitz.open(str(path))
    except Exception as e:
        print(f"[PDF PARSER ERROR] Failed to open PDF '{path}': {e}")
        raise

    total_pages = len(doc)
    print(f"[PDF PARSER] Document opened successfully. Total pages: {total_pages}")

    pages: list[str] = []
    for page_num in range(total_pages):
        page = doc.load_page(page_num)
        text = page.get_text()
        if text.strip():
            pages.append(f"[Page {page_num + 1}]\n{text}")

    full_text = "\n\n".join(pages)
    print(f"[PDF PARSER] Processed {total_pages} pages, extracted {len(full_text):,} total chars")

    if len(full_text) < 50:
        print("[PDF PARSER ERROR] Extracted text length < 50 chars. Suspected scanned image PDF.")
        raise ValueError(
            "PDF appears to be scanned or image-based. "
            "Please upload digital text PDFs for the EduNexus pilot."
        )

    return full_text


def _parse_pptx(path: Path) -> str:
    print(f"[PPTX PARSER] Parsing presentation: {path}")

    try:
        from pptx import Presentation  # type: ignore[import-untyped]
    except ImportError as e:
        print(f"[PPTX PARSER ERROR] python-pptx is not installed: {e}")
        raise RuntimeError("python-pptx is required. Install via `pip install python-pptx`.") from e

    try:
        prs = Presentation(str(path))
    except Exception as e:
        print(f"[PPTX PARSER ERROR] Failed to open presentation '{path}': {e}")
        raise

    total_slides = len(prs.slides)
    print(f"[PPTX PARSER] Presentation opened. Total slides: {total_slides}")

    slides: list[str] = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if line:
                        slide_lines.append(line)
        if slide_lines:
            slides.append(f"[Slide {slide_num}]\n" + "\n".join(slide_lines))

    combined = "\n\n".join(slides)
    print(f"[PPTX PARSER] Processed {total_slides} slides, extracted {len(combined):,} total chars")
    return combined


def _parse_docx(path: Path) -> str:
    print(f"[DOCX PARSER] Parsing Word document: {path}")

    try:
        from docx import Document  # type: ignore[import-untyped]
    except ImportError as e:
        print(f"[DOCX PARSER ERROR] python-docx is not installed: {e}")
        raise RuntimeError("python-docx is required. Install via `pip install python-docx`.") from e

    try:
        doc = Document(str(path))
    except Exception as e:
        print(f"[DOCX PARSER ERROR] Failed to open document '{path}': {e}")
        raise

    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    combined = "\n".join(paragraphs)
    print(f"[DOCX PARSER] Processed {len(paragraphs)} paragraphs, extracted {len(combined):,} total chars")
    return combined


def _parse_txt(path: Path) -> str:
    print(f"[TXT PARSER] Reading plain text file: {path}")

    try:
        text = path.read_text(encoding="utf-8")
        print(f"[TXT PARSER] Read {len(text):,} chars with UTF-8 encoding")
    except UnicodeDecodeError:
        print(f"[TXT PARSER] UTF-8 decode failed for '{path.name}'. Retrying with latin-1.")
        logger.warning("UTF-8 decode failed for '%s'; retrying with latin-1.", path.name)
        try:
            text = path.read_text(encoding="latin-1")
            print(f"[TXT PARSER] Read {len(text):,} chars with latin-1 encoding")
        except Exception as e:
            print(f"[TXT PARSER ERROR] Failed to read file '{path}': {e}")
            raise

    return text
